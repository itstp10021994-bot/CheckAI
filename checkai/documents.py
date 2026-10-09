"""Trích văn bản từ tệp Word (.docx), PDF và văn bản thuần (.txt, .md)."""

from __future__ import annotations

import io
import re

MAX_PDF_PAGES = 300
SUPPORTED = (".docx", ".pdf", ".txt", ".md")
_END = re.compile(r"[.!?…:;\"”)»]\s*$")


def _docx(data: bytes) -> tuple[str, dict]:
    import docx  # python-docx

    try:
        doc = docx.Document(io.BytesIO(data))
    except Exception as e:  # noqa: BLE001
        raise ValueError("Không đọc được tệp Word. Hãy chắc chắn tệp là .docx hợp lệ.") from e
    paras = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                paras.append(" | ".join(dict.fromkeys(cells)))  # bỏ ô gộp bị lặp
    cp = doc.core_properties
    meta = {k: str(v) for k, v in {
        "Tác giả": cp.author, "Sửa lần cuối bởi": cp.last_modified_by,
        "Ngày tạo": cp.created, "Ngày sửa": cp.modified,
    }.items() if v}
    return "\n\n".join(paras), meta


def _join_pdf_lines(page: str) -> str:
    """PDF ngắt dòng cứng theo chiều rộng trang; nối lại thành đoạn văn để tách câu đúng."""
    lines = [ln.strip() for ln in page.splitlines()]
    width = max((len(ln) for ln in lines), default=0)
    paras: list[str] = []
    cur = ""
    for ln in lines:
        if not ln:
            if cur:
                paras.append(cur)
                cur = ""
            continue
        if cur.endswith("-") and not cur.endswith(" -"):
            cur = cur[:-1] + ln  # từ bị ngắt bằng gạch nối
        else:
            cur = f"{cur} {ln}".strip()
        # Dòng kết thúc câu và ngắn hơn hẳn chiều rộng trang -> hết đoạn.
        if _END.search(ln) and len(ln) < 0.8 * width:
            paras.append(cur)
            cur = ""
    if cur:
        paras.append(cur)
    return "\n\n".join(paras)


def _pdf(data: bytes) -> tuple[str, dict]:
    from pypdf import PdfReader

    try:
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted and not reader.decrypt(""):
            raise ValueError("PDF được đặt mật khẩu, không đọc được.")
        pages = reader.pages[:MAX_PDF_PAGES]
        text = "\n\n".join(_join_pdf_lines(p.extract_text() or "") for p in pages)
    except ValueError:
        raise
    except Exception as e:  # noqa: BLE001
        raise ValueError("Không đọc được tệp PDF.") from e
    meta = {"Số trang": str(len(reader.pages))}
    info = reader.metadata or {}
    for key, label in (("/Author", "Tác giả"), ("/Creator", "Phần mềm tạo"), ("/Producer", "Phần mềm xuất PDF")):
        if info.get(key):
            meta[label] = str(info[key])
    if len(reader.pages) > MAX_PDF_PAGES:
        meta["Ghi chú"] = f"Chỉ đọc {MAX_PDF_PAGES} trang đầu."
    if len(re.findall(r"\w+", text)) < 5:
        raise ValueError("PDF không có lớp chữ (có thể là ảnh scan). Hiện chưa hỗ trợ nhận dạng chữ (OCR).")
    return text, meta


def extract_text(data: bytes, filename: str) -> tuple[str, dict]:
    """Trả về (văn bản, thông tin tệp). Ném ValueError nếu không hỗ trợ / không đọc được."""
    name = (filename or "").lower()
    if name.endswith(".docx"):
        text, meta = _docx(data)
    elif name.endswith(".pdf"):
        text, meta = _pdf(data)
    elif name.endswith((".txt", ".md")):
        text, meta = data.decode("utf-8", errors="ignore"), {}
    elif name.endswith(".doc"):
        raise ValueError("Định dạng .doc (Word 97-2003) chưa được hỗ trợ. Hãy mở bằng Word và lưu lại thành .docx.")
    else:
        raise ValueError("Chỉ hỗ trợ tệp " + ", ".join(SUPPORTED) + ".")
    text = text.strip()
    if not text:
        raise ValueError("Tệp không có nội dung văn bản.")
    return text, {"Tệp": filename, **meta, "Số ký tự": str(len(text))}
